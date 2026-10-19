---

### 3. `app.py`
```python
import io
import re
import sqlite3
import random
from typing import List, Dict, Any, Tuple
from PIL import Image

import streamlit as st

# Document Parsing Libraries
import pypdf
import pdfplumber
import docx

# PDF Generation Libraries (ReportLab Platypus)
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

# Word Generation Libraries
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ==========================================
# 1. DATABASE INITIALIZATION & OPERATIONS
# ==========================================

DB_FILE = "question_bank.db"

def init_db():
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS question_bank (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            subject TEXT NOT NULL,
            chapter TEXT,
            question_text TEXT NOT NULL,
            is_mcq INTEGER NOT NULL,
            option_a TEXT,
            option_b TEXT,
            option_c TEXT,
            option_d TEXT,
            correct_answer TEXT,
            marks INTEGER NOT NULL,
            difficulty TEXT,
            blooms_level TEXT
        )
    """)
    conn.commit()
    conn.close()

def save_question_to_db(q_data: Dict[str, Any]):
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO question_bank 
        (subject, chapter, question_text, is_mcq, option_a, option_b, option_c, option_d, correct_answer, marks, difficulty, blooms_level)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        q_data.get("subject", "General"),
        q_data.get("chapter", "General"),
        q_data["text"],
        1 if q_data["is_mcq"] else 0,
        q_data["options"].get("A", "") if q_data["is_mcq"] else "",
        q_data["options"].get("B", "") if q_data["is_mcq"] else "",
        q_data["options"].get("C", "") if q_data["is_mcq"] else "",
        q_data["options"].get("D", "") if q_data["is_mcq"] else "",
        q_data.get("correct_answer", ""),
        q_data.get("marks", 1),
        q_data.get("difficulty", "Medium"),
        q_data.get("blooms_level", "Understand")
    ))
    conn.commit()
    conn.close()

def get_filtered_questions(subject: str = None, difficulty: str = None) -> List[Dict[str, Any]]:
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    query = "SELECT id, subject, chapter, question_text, is_mcq, option_a, option_b, option_c, option_d, correct_answer, marks, difficulty, blooms_level FROM question_bank WHERE 1=1"
    params = []
    if subject and subject != "All":
        query += " AND subject = ?"
        params.append(subject)
    if difficulty and difficulty != "All":
        query += " AND difficulty = ?"
        params.append(difficulty)
    
    cursor.execute(query, params)
    rows = cursor.fetchall()
    conn.close()
    
    questions = []
    for r in rows:
        questions.append({
            "id": r[0],
            "subject": r[1],
            "chapter": r[2],
            "text": r[3],
            "is_mcq": bool(r[4]),
            "options": {"A": r[5], "B": r[6], "C": r[7], "D": r[8]} if r[4] else {"A": "", "B": "", "C": "", "D": ""},
            "correct_answer": r[9],
            "marks": r[10],
            "difficulty": r[11],
            "blooms_level": r[12]
        })
    return questions

def get_unique_subjects() -> List[str]:
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    cursor.execute("SELECT DISTINCT subject FROM question_bank")
    rows = cursor.fetchall()
    conn.close()
    return [r[0] for r in rows if r[0]]

# ==========================================
# 2. APPLICATION STATE INITIALIZATION
# ==========================================

def init_session_state():
    if "metadata" not in st.session_state:
        st.session_state.metadata = {
            "institution_name": "ROYAL INSTITUTE OF TECHNOLOGY & SCIENCE",
            "department": "DEPARTMENT OF COMPUTER SCIENCE & ENGINEERING",
            "exam_name": "END SEMESTER EXAMINATION - FALL 2026",
            "subject_title": "ADVANCED ARTIFICIAL INTELLIGENCE",
            "subject_code": "CS-8001",
            "date": "2026-11-15",
            "duration": "3 Hours",
            "max_marks": 100,
            "instructions": (
                "1. All questions in Section A are compulsory.\n"
                "2. Answer any four questions from Section B.\n"
                "3. Figures to the right indicate full marks.\n"
                "4. Use of programmable calculators is strictly prohibited."
            ),
            "logo_bytes": None
        }
    if "sections" not in st.session_state:
        st.session_state.sections = [
            {
                "id": "sec_a",
                "title": "Section A: Multiple Choice Questions",
                "is_mcq": True,
                "marks_per_q": 2,
                "questions": [
                    {
                        "text": "Which search algorithm guarantees finding the optimal path if the heuristic is admissible?",
                        "is_mcq": True,
                        "options": {
                            "A": "Breadth-First Search",
                            "B": "A* Search",
                            "C": "Depth-First Search",
                            "D": "Greedy Best-First Search"
                        },
                        "correct_answer": "B",
                        "difficulty": "Easy",
                        "blooms_level": "Remember"
                    },
                    {
                        "text": "In deep learning, which activation function mitigates the vanishing gradient problem?",
                        "is_mcq": True,
                        "options": {
                            "A": "Sigmoid",
                            "B": "Tanh",
                            "C": "ReLU",
                            "D": "Binary Step"
                        },
                        "correct_answer": "C",
                        "difficulty": "Medium",
                        "blooms_level": "Understand"
                    }
                ]
            },
            {
                "id": "sec_b",
                "title": "Section B: Short Answer & Analytical Questions",
                "is_mcq": False,
                "marks_per_q": 10,
                "questions": [
                    {
                        "text": "Explain the architectural differences between Transformer models and traditional Recurrent Neural Networks (RNNs). Highlight self-attention mechanisms.",
                        "is_mcq": False,
                        "options": {"A": "", "B": "", "C": "", "D": ""},
                        "correct_answer": "RNNs process sequentially creating bottlenecks, whereas Transformers process tokens in parallel using Self-Attention vectors (Query, Key, Value) to capture long-range dependencies.",
                        "difficulty": "Hard",
                        "blooms_level": "Analyze"
                    }
                ]
            }
        ]

# ==========================================
# 3. TEXT & DOCUMENT PARSING ENGINE
# ==========================================

def parse_raw_text_to_questions(text: str) -> List[Dict[str, Any]]:
    questions = []
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    
    current_q = None
    mcq_pattern = re.compile(r'^\(?([A-Da-d])[\)\.]\s*(.*)')
    q_pattern = re.compile(r'^(?:Q|Question)?\s*\d+[\.\)]\s*(.*)', re.IGNORECASE)

    for line in lines:
        q_match = q_pattern.match(line)
        if q_match:
            if current_q:
                questions.append(current_q)
            q_text = q_match.group(1).strip()
            current_q = {
                "text": q_text,
                "is_mcq": False,
                "options": {"A": "", "B": "", "C": "", "D": ""},
                "correct_answer": "",
                "difficulty": "Medium",
                "blooms_level": "Understand"
            }
            continue

        opt_match = mcq_pattern.match(line)
        if opt_match and current_q:
            opt_key = opt_match.group(1).upper()
            opt_val = opt_match.group(2).strip()
            current_q["is_mcq"] = True
            current_q["options"][opt_key] = opt_val
        elif current_q and not current_q["is_mcq"]:
            current_q["text"] += " " + line

    if current_q:
        questions.append(current_q)
        
    return questions

def extract_text_from_pdf(file_bytes: bytes) -> str:
    extracted_text = ""
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    extracted_text += text + "\n"
    except Exception:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            extracted_text += (page.extract_text() or "") + "\n"
    return extracted_text

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    return "\n".join([p.text for p in doc.paragraphs if p.text])

# ==========================================
# 4. REPORTLAB PDF GENERATION ENGINE
# ==========================================

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setStrokeColor(colors.HexColor("#1e293b"))
        self.setLineWidth(1)
        
        # Outer Double Border
        self.rect(20, 20, 572, 752)
        self.rect(23, 23, 566, 746)

        # Footer Page Number
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#475569"))
        footer_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(560, 32, footer_text)
        self.drawString(32, 32, "CONFIDENTIAL - EXAMINATION PAPER")
        self.restoreState()

def generate_pdf_paper(metadata: Dict[str, Any], sections: List[Dict[str, Any]], is_answer_key: bool = False) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()
    c_primary = colors.HexColor("#0f172a")
    c_secondary = colors.HexColor("#334155")
    
    style_inst = ParagraphStyle('InstHeader', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=14, leading=16, alignment=1, textColor=c_primary)
    style_dept = ParagraphStyle('DeptHeader', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, leading=12, alignment=1, textColor=c_secondary)
    style_exam = ParagraphStyle('ExamHeader', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=13, alignment=1, textColor=c_primary)
    style_meta = ParagraphStyle('MetaText', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=11, textColor=c_primary)
    style_meta_bold = ParagraphStyle('MetaTextBold', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9, leading=11, textColor=c_primary)
    style_sec = ParagraphStyle('SecHeader', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=13, textColor=c_primary)
    style_q_text = ParagraphStyle('QText', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=13, textColor=colors.black)
    style_opt = ParagraphStyle('OptText', parent=styles['Normal'], fontName='Helvetica', fontSize=9.5, leading=12, textColor=colors.HexColor("#1e293b"))
    style_ans = ParagraphStyle('AnsText', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9.5, leading=12, textColor=colors.HexColor("#166534"))

    story = []

    logo_img = None
    if metadata.get("logo_bytes"):
        try:
            img_data = io.BytesIO(metadata["logo_bytes"])
            logo_img = RLImage(img_data, width=60, height=60)
        except Exception:
            logo_img = None

    title_p = [
        Paragraph(metadata["institution_name"].upper(), style_inst),
        Spacer(1, 2),
        Paragraph(metadata["department"].upper(), style_dept),
        Spacer(1, 2),
        Paragraph(f"<b>{metadata['exam_name']}</b>" + (" - ANSWER KEY & RUBRIC" if is_answer_key else ""), style_exam)
    ]

    if logo_img:
        header_table = Table([[logo_img, title_p]], colWidths=[70, 470])
        header_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('ALIGN', (0,0), (0,0), 'CENTER')]))
    else:
        header_table = Table([[title_p]], colWidths=[540])
        header_table.setStyle(TableStyle([('ALIGN', (0,0), (-1,-1), 'CENTER')]))

    story.append(header_table)
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1.5, color=c_primary, spaceBefore=2, spaceAfter=8))

    meta_table_data = [
        [
            Paragraph(f"<b>Course / Subject:</b> {metadata['subject_title']}", style_meta),
            Paragraph(f"<b>Subject Code:</b> {metadata['subject_code']}", style_meta)
        ],
        [
            Paragraph(f"<b>Date of Exam:</b> {metadata['date']}", style_meta),
            Paragraph(f"<b>Duration:</b> {metadata['duration']}", style_meta)
        ],
        [
            Paragraph(f"<b>Max Marks:</b> {metadata['max_marks']}", style_meta_bold),
            Paragraph(f"<b>Academic Year:</b> 2026-2027", style_meta)
        ]
    ]
    meta_table = Table(meta_table_data, colWidths=[330, 210])
    meta_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('BOTTOMPADDING', (0,0), (-1,-1), 2), ('TOPPADDING', (0,0), (-1,-1), 2)]))
    story.append(meta_table)
    story.append(Spacer(1, 6))
    story.append(HRFlowable(width="100%", thickness=0.5, color=c_secondary, spaceBefore=2, spaceAfter=8))

    if metadata.get("instructions") and not is_answer_key:
        story.append(Paragraph("<b>General Instructions:</b>", style_meta_bold))
        story.append(Spacer(1, 2))
        for line in metadata["instructions"].splitlines():
            if line.strip():
                story.append(Paragraph(f"• {line.strip()}", style_meta))
                story.append(Spacer(1, 1))
        story.append(Spacer(1, 8))
        story.append(HRFlowable(width="100%", thickness=0.5, color=c_secondary, spaceBefore=2, spaceAfter=10))

    q_counter = 1
    for sec in sections:
        total_sec_q = len(sec["questions"])
        sec_marks_str = f"[{total_sec_q} × {sec['marks_per_q']} = {total_sec_q * sec['marks_per_q']} Marks]"
        sec_table_data = [
            [Paragraph(sec["title"].upper(), style_sec), Paragraph(f"<b>{sec_marks_str}</b>", ParagraphStyle('RMark', parent=style_sec, alignment=2))]
        ]
        sec_table = Table(sec_table_data, colWidths=[400, 140])
        sec_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('BOTTOMPADDING', (0,0), (-1,-1), 2), ('TOPPADDING', (0,0), (-1,-1), 2)]))
        story.append(sec_table)
        story.append(Spacer(1, 6))

        for q in sec["questions"]:
            q_str = f"<b>Q{q_counter}.</b> {q['text']}"
            q_mark_str = f"[{sec['marks_per_q']} Mark{'s' if sec['marks_per_q'] > 1 else ''}]"
            
            q_row_table = Table([
                [Paragraph(q_str, style_q_text), Paragraph(q_mark_str, ParagraphStyle('QMark', parent=style_meta_bold, alignment=2))]
            ], colWidths=[470, 70])
            q_row_table.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('LEFTPADDING', (0,0), (-1,-1), 0), ('RIGHTPADDING', (0,0), (-1,-1), 0)]))
            story.append(q_row_table)
            story.append(Spacer(1, 4))

            if sec["is_mcq"] or q.get("is_mcq"):
                opts = q.get("options", {})
                opt_a = Paragraph(f"<b>(A)</b> {opts.get('A', '')}", style_opt)
                opt_b = Paragraph(f"<b>(B)</b> {opts.get('B', '')}", style_opt)
                opt_c = Paragraph(f"<b>(C)</b> {opts.get('C', '')}", style_opt)
                opt_d = Paragraph(f"<b>(D)</b> {opts.get('D', '')}", style_opt)

                opts_grid = Table([[opt_a, opt_b], [opt_c, opt_d]], colWidths=[260, 260])
                opts_grid.setStyle(TableStyle([('VALIGN', (0,0), (-1,-1), 'TOP'), ('LEFTPADDING', (0,0), (-1,-1), 15), ('BOTTOMPADDING', (0,0), (-1,-1), 3), ('TOPPADDING', (0,0), (-1,-1), 2)]))
                story.append(opts_grid)
                story.append(Spacer(1, 4))

            if is_answer_key and q.get("correct_answer"):
                ans_str = f"<b>Correct Answer / Rubric:</b> Option ({q['correct_answer']})" if (sec["is_mcq"] or q.get("is_mcq")) else f"<b>Key Solution / Rubric:</b> {q['correct_answer']}"
                ans_table = Table([[Paragraph(ans_str, style_ans)]], colWidths=[540])
                ans_table.setStyle(TableStyle([('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#f0fdf4")), ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#bbf7d0")), ('LEFTPADDING', (0,0), (-1,-1), 8), ('RIGHTPADDING', (0,0), (-1,-1), 8), ('TOPPADDING', (0,0), (-1,-1), 4), ('BOTTOMPADDING', (0,0), (-1,-1), 4)]))
                story.append(Spacer(1, 2))
                story.append(ans_table)
                story.append(Spacer(1, 4))

            story.append(Spacer(1, 4))
            q_counter += 1

        story.append(Spacer(1, 8))

    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()

# ==========================================
# 5. PYTHON-DOCX WORD GENERATION ENGINE
# ==========================================

def generate_docx_paper(metadata: Dict[str, Any], sections: List[Dict[str, Any]], is_answer_key: bool = False) -> bytes:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(10)
    style_normal.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    if metadata.get("logo_bytes"):
        header_table = doc.add_table(rows=1, cols=2)
        header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        header_table.autofit = False
        
        cell_logo = header_table.cell(0, 0)
        cell_logo.width = Inches(1.2)
        img_stream = io.BytesIO(metadata["logo_bytes"])
        p_img = cell_logo.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(img_stream, width=Inches(1.0))

        cell_text = header_table.cell(0, 1)
        cell_text.width = Inches(6.3)
    else:
        header_table = doc.add_table(rows=1, cols=1)
        header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell_text = header_table.cell(0, 0)
        cell_text.width = Inches(7.5)

    p_inst = cell_text.paragraphs[0]
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run(metadata["institution_name"].upper() + "\n")
    r_inst.bold = True
    r_inst.font.size = Pt(14)
    r_inst.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    r_dept = p_inst.add_run(metadata["department"].upper() + "\n")
    r_dept.bold = True
    r_dept.font.size = Pt(10)
    r_dept.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    r_exam = p_inst.add_run(metadata["exam_name"].upper() + (" - ANSWER KEY" if is_answer_key else ""))
    r_exam.bold = True
    r_exam.font.size = Pt(11)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    grid_data = [
        (f"Course / Subject: {metadata['subject_title']}", f"Subject Code: {metadata['subject_code']}"),
        (f"Date of Exam: {metadata['date']}", f"Duration: {metadata['duration']}"),
        (f"Max Marks: {metadata['max_marks']}", "Academic Year: 2026-2027")
    ]

    for r_idx, row_content in enumerate(grid_data):
        row = meta_table.rows[r_idx]
        cell_l, cell_r = row.cells[0], row.cells[1]
        cell_l.width, cell_r.width = Inches(4.5), Inches(3.0)
        
        p_l = cell_l.paragraphs[0]
        p_l.add_run(row_content[0]).bold = ("Max Marks" in row_content[0])
        p_l.paragraph_format.space_after = Pt(2)
        
        p_r = cell_r.paragraphs[0]
        p_r.add_run(row_content[1])
        p_r.paragraph_format.space_after = Pt(2)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    if metadata.get("instructions") and not is_answer_key:
        p_inst_h = doc.add_paragraph()
        r_inst_h = p_inst_h.add_run("General Instructions:")
        r_inst_h.bold = True
        p_inst_h.paragraph_format.space_after = Pt(2)

        for line in metadata["instructions"].splitlines():
            if line.strip():
                p_i = doc.add_paragraph(style='List Bullet')
                p_i.add_run(line.strip())
                p_i.paragraph_format.space_after = Pt(1)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    q_counter = 1
    for sec in sections:
        sec_table = doc.add_table(rows=1, cols=2)
        sec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sec_cell_l, sec_cell_r = sec_table.cell(0, 0), sec_table.cell(0, 1)
        sec_cell_l.width, sec_cell_r.width = Inches(5.5), Inches(2.0)

        p_sec_l = sec_cell_l.paragraphs[0]
        r_sec_title = p_sec_l.add_run(sec["title"].upper())
        r_sec_title.bold = True
        r_sec_title.font.size = Pt(11)

        total_sec_q = len(sec["questions"])
        sec_marks_str = f"[{total_sec_q} × {sec['marks_per_q']} = {total_sec_q * sec['marks_per_q']} Marks]"
        p_sec_r = sec_cell_r.paragraphs[0]
        p_sec_r.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_sec_marks = p_sec_r.add_run(sec_marks_str)
        r_sec_marks.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        for q in sec["questions"]:
            q_table = doc.add_table(rows=1, cols=2)
            q_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            q_cell_l, q_cell_r = q_table.cell(0, 0), q_table.cell(0, 1)
            q_cell_l.width, q_cell_r.width = Inches(6.5), Inches(1.0)

            p_q = q_cell_l.paragraphs[0]
            r_q_num = p_q.add_run(f"Q{q_counter}. ")
            r_q_num.bold = True
            p_q.add_run(q["text"])

            p_m = q_cell_r.paragraphs[0]
            p_m.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_m.add_run(f"[{sec['marks_per_q']} Mark{'s' if sec['marks_per_q'] > 1 else ''}]").bold = True

            if sec["is_mcq"] or q.get("is_mcq"):
                opts = q.get("options", {})
                opts_table = doc.add_table(rows=2, cols=2)
                opts_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                opts_grid = [
                    [("A", opts.get("A", "")), ("B", opts.get("B", ""))],
                    [("C", opts.get("C", "")), ("D", opts.get("D", ""))]
                ]
                for row_i, row_opts in enumerate(opts_grid):
                    for col_i, (opt_k, opt_v) in enumerate(row_opts):
                        c = opts_table.cell(row_i, col_i)
                        c.width = Inches(3.7)
                        p_opt = c.paragraphs[0]
                        p_opt.add_run(f"({opt_k}) ").bold = True
                        p_opt.add_run(opt_v)
                        p_opt.paragraph_format.space_after = Pt(2)

            if is_answer_key and q.get("correct_answer"):
                p_ans = doc.add_paragraph()
                p_ans.paragraph_format.space_before = Pt(2)
                p_ans.paragraph_format.space_after = Pt(4)
                r_ans = p_ans.add_run(f"Key Solution/Rubric: {q['correct_answer']}")
                r_ans.bold = True
                r_ans.font.color.rgb = RGBColor(0x16, 0x65, 0x34)

            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            q_counter += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

# ==========================================
# 6. HIGH-END GLASSMORPHISM CSS INJECTION
# ==========================================

def inject_glassmorphism_css():
    css = """
    <style>
    @import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600;700;800&display=swap');

    html, body, [data-testid="stAppViewContainer"] {
        font-family: 'Plus Jakarta Sans', sans-serif !important;
        background: linear-gradient(135deg, #0f172a 0%, #1e1b4b 50%, #0f172a 100%) !important;
        background-attachment: fixed !important;
        color: #f8fafc !important;
    }

    [data-testid="stHeader"] {
        background: rgba(15, 23, 42, 0.6) !important;
        backdrop-filter: blur(12px) !important;
    }

    [data-testid="stSidebar"] {
        background: rgba(15, 23, 42, 0.75) !important;
        backdrop-filter: blur(20px) !important;
        border-right: 1px solid rgba(255, 255, 255, 0.1) !important;
    }

    div.glass-card {
        background: rgba(255, 255, 255, 0.04) !important;
        backdrop-filter: blur(16px) !important;
        -webkit-backdrop-filter: blur(16px) !important;
        border: 1px solid rgba(255, 255, 255, 0.12) !important;
        border-radius: 16px !important;
        padding: 24px !important;
        margin-bottom: 20px !important;
        box-shadow: 0 8px 32px 0 rgba(0, 0, 0, 0.37) !important;
    }

    div.glass-card-accent {
        background: linear-gradient(135deg, rgba(99, 102, 241, 0.1) 0%, rgba(139, 92, 246, 0.1) 100%) !important;
        backdrop-filter: blur(16px) !important;
        border: 1px solid rgba(99, 102, 241, 0.3) !important;
        border-radius: 16px !important;
        padding: 20px !important;
        margin-bottom: 16px !important;
    }

    .stTextInput > div > div > input,
    .stTextArea > div > div > textarea,
    .stSelectbox > div > div > div,
    .stNumberInput > div > div > input {
        background: rgba(15, 23, 42, 0.6) !important;
        color: #f8fafc !important;
        border: 1px solid rgba(255, 255, 255, 0.15) !important;
        border-radius: 10px !important;
        backdrop-filter: blur(10px) !important;
    }

    .stButton > button {
        background: linear-gradient(135deg, #6366f1 0%, #8b5cf6 100%) !important;
        color: #ffffff !important;
        font-weight: 600 !important;
        border: none !important;
        border-radius: 10px !important;
        padding: 10px 22px !important;
        transition: all 0.3s ease !important;
        box-shadow: 0 4px 15px rgba(99, 102, 241, 0.35) !important;
    }

    .stButton > button:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 6px 20px rgba(139, 92, 246, 0.5) !important;
    }

    .stTabs [data-baseweb="tab-list"] {
        gap: 8px !important;
        background: rgba(15, 23, 42, 0.5) !important;
        padding: 8px !important;
        border-radius: 12px !important;
        border: 1px solid rgba(255, 255, 255, 0.08) !important;
    }

    .stTabs [data-baseweb="tab"] {
        height: 44px !important;
        border-radius: 8px !important;
        color: #94a3b8 !important;
        font-weight: 600 !important;
        border: none !important;
    }

    .stTabs [aria-selected="true"] {
        background: linear-gradient(135deg, #06b6d4 0%, #3b82f6 100%) !important;
        color: #ffffff !important;
    }

    div.metric-card {
        background: rgba(255, 255, 255, 0.03);
        border: 1px solid rgba(255, 255, 255, 0.1);
        border-radius: 12px;
        padding: 16px;
        text-align: center;
    }
    div.metric-val {
        font-size: 24px;
        font-weight: 800;
        color: #06b6d4;
    }
    div.metric-lbl {
        font-size: 12px;
        color: #94a3b8;
        text-transform: uppercase;
        letter-spacing: 1px;
    }
    </style>
    """
    st.markdown(css, unsafe_allow_html=True)

# ==========================================
# 7. MAIN STREAMLIT APPLICATION INTERFACE
# ==========================================

def main():
    st.set_page_config(
        page_title="PaperCraft AI | Professional Exam Generator",
        page_icon="🎓",
        layout="wide",
        initial_sidebar_state="expanded"
    )

    init_db()
    init_session_state()
    inject_glassmorphism_css()

    st.markdown("""
        <div class="glass-card" style="display: flex; align-items: center; justify-content: space-between; border-left: 6px solid #6366f1;">
            <div>
                <h1 style="margin: 0; font-size: 32px; font-weight: 800; background: linear-gradient(90deg, #38bdf8, #818cf8, #c084fc); -webkit-background-clip: text; -webkit-text-fill-color: transparent;">
                    PaperCraft AI
                </h1>
                <p style="margin: 4px 0 0 0; color: #94a3b8; font-size: 15px;">
                    Enterprise Examination Paper Engineering & Multi-Format Publishing System
                </p>
            </div>
            <div>
                <span style="background: rgba(99, 102, 241, 0.2); border: 1px solid #6366f1; color: #a5b4fc; padding: 6px 14px; border-radius: 20px; font-size: 13px; font-weight: 600;">
                    v2.5 Production Build
                </span>
            </div>
        </div>
    """, unsafe_allow_html=True)

    total_calculated_marks = sum(
        len(sec["questions"]) * sec["marks_per_q"] for sec in st.session_state.sections
    )
    total_questions_count = sum(len(sec["questions"]) for sec in st.session_state.sections)

    m1, m2, m3, m4 = st.columns(4)
    with m1:
        st.markdown(f'<div class="metric-card"><div class="metric-val">{total_questions_count}</div><div class="metric-lbl">Total Questions</div></div>', unsafe_allow_html=True)
    with m2:
        st.markdown(f'<div class="metric-card"><div class="metric-val">{len(st.session_state.sections)}</div><div class="metric-lbl">Active Sections</div></div>', unsafe_allow_html=True)
    with m3:
        st.markdown(f'<div class="metric-card"><div class="metric-val">{total_calculated_marks} / {st.session_state.metadata["max_marks"]}</div><div class="metric-lbl">Allocated Marks</div></div>', unsafe_allow_html=True)
    with m4:
        status_color = "#10b981" if total_calculated_marks == st.session_state.metadata["max_marks"] else "#f59e0b"
        status_text = "Balanced" if total_calculated_marks == st.session_state.metadata["max_marks"] else "Mismatch"
        st.markdown(f'<div class="metric-card"><div class="metric-val" style="color: {status_color};">{status_text}</div><div class="metric-lbl">Marks Validation</div></div>', unsafe_allow_html=True)

    st.markdown("<br>", unsafe_allow_html=True)

    with st.sidebar:
        st.markdown("<h2 style='font-size: 18px; color: #f8fafc;'>🏛️ Institutional Branding</h2>", unsafe_allow_html=True)
        logo_file = st.file_uploader("Upload Institution Logo (PNG/JPG)", type=["png", "jpg", "jpeg"])
        if logo_file:
            st.session_state.metadata["logo_bytes"] = logo_file.read()
            st.image(st.session_state.metadata["logo_bytes"], width=120)

        st.markdown("<hr style='border-color: rgba(255,255,255,0.1);'>", unsafe_allow_html=True)
        st.markdown("<h2 style='font-size: 18px; color: #f8fafc;'>📋 Exam Metadata Header</h2>", unsafe_allow_html=True)

        meta = st.session_state.metadata
        meta["institution_name"] = st.text_input("Institution Name", value=meta["institution_name"])
        meta["department"] = st.text_input("Department Title", value=meta["department"])
        meta["exam_name"] = st.text_input("Examination Name", value=meta["exam_name"])
        meta["subject_title"] = st.text_input("Course / Subject Name", value=meta["subject_title"])
        meta["subject_code"] = st.text_input("Subject Code", value=meta["subject_code"])
        
        c_d1, c_d2 = st.columns(2)
        with c_d1:
            meta["date"] = st.text_input("Date", value=meta["date"])
        with c_d2:
            meta["duration"] = st.text_input("Duration", value=meta["duration"])
            
        meta["max_marks"] = st.number_input("Maximum Target Marks", value=meta["max_marks"], step=5)
        meta["instructions"] = st.text_area("General Instructions (1 per line)", value=meta["instructions"], height=120)

    tab_builder, tab_parser, tab_bank, tab_randomizer, tab_export = st.tabs([
        "📝 Interactive Paper Builder",
        "📂 Sample Paper Parser",
        "🗄️ Question Bank Repository",
        "🔀 Multi-Set Randomizer",
        "🚀 High-Precision Export Engine"
    ])

    with tab_builder:
        st.markdown("<div class='glass-card'>", unsafe_allow_html=True)
        st.subheader("Dynamic Section Management")
        
        col_sec_btn, _ = st.columns([1, 3])
        with col_sec_btn:
            if st.button("➕ Add New Section"):
                next_letter = chr(65 + len(st.session_state.sections))
                st.session_state.sections.append({
                    "id": f"sec_{next_letter.lower()}_{random.randint(100,999)}",
                    "title": f"Section {next_letter}: Short Descriptive Questions",
                    "is_mcq": False,
                    "marks_per_q": 5,
                    "questions": []
                })
                st.rerun()
        st.markdown("</div>", unsafe_allow_html=True)

        sec_to_remove = None
        for s_idx, sec in enumerate(st.session_state.sections):
            with st.expander(f"📌 {sec['title']} ({len(sec['questions'])} Questions | {sec['marks_per_q']} Marks Each)", expanded=True):
                st.markdown("<div class='glass-card-accent'>", unsafe_allow_html=True)
                sc1, sc2, sc3 = st.columns([3, 1, 1])
                with sc1:
                    sec["title"] = st.text_input("Section Header Text", value=sec["title"], key=f"sec_title_{sec['id']}")
                with sc2:
                    sec["marks_per_q"] = st.number_input("Marks / Question", value=sec["marks_per_q"], min_value=1, key=f"sec_marks_{sec['id']}")
                with sc3:
                    sec["is_mcq"] = st.checkbox("Is MCQ Section?", value=sec["is_mcq"], key=f"sec_mcq_{sec['id']}")
                
                if st.button("🗑️ Delete Section", key=f"del_sec_{sec['id']}"):
                    sec_to_remove = s_idx
                st.markdown("</div>", unsafe_allow_html=True)

                q_to_remove = None
                for q_idx, q in enumerate(sec["questions"]):
                    st.markdown(f"**Question {q_idx + 1}**")
                    q_col1, q_col2 = st.columns([5, 1])
                    
                    with q_col1:
                        q["text"] = st.text_area("Question Stem", value=q["text"], key=f"q_text_{sec['id']}_{q_idx}", height=70)
                        
                        if sec["is_mcq"]:
                            st.caption("Multiple Choice Options")
                            o1, o2 = st.columns(2)
                            with o1:
                                q["options"]["A"] = st.text_input("Option A", value=q["options"].get("A", ""), key=f"opt_a_{sec['id']}_{q_idx}")
                                q["options"]["C"] = st.text_input("Option C", value=q["options"].get("C", ""), key=f"opt_c_{sec['id']}_{q_idx}")
                            with o2:
                                q["options"]["B"] = st.text_input("Option B", value=q["options"].get("B", ""), key=f"opt_b_{sec['id']}_{q_idx}")
                                q["options"]["D"] = st.text_input("Option D", value=q["options"].get("D", ""), key=f"opt_d_{sec['id']}_{q_idx}")
                            
                            q["correct_answer"] = st.selectbox("Correct Option Key", ["A", "B", "C", "D"], index=["A", "B", "C", "D"].index(q.get("correct_answer", "A")) if q.get("correct_answer") in ["A", "B", "C", "D"] else 0, key=f"corr_{sec['id']}_{q_idx}")
                        else:
                            q["correct_answer"] = st.text_area("Key Solution / Marking Rubric", value=q.get("correct_answer", ""), key=f"rubric_{sec['id']}_{q_idx}", height=60)

                    with q_col2:
                        st.markdown("<br>", unsafe_allow_html=True)
                        btn_up, btn_dn = st.columns(2)
                        with btn_up:
                            if st.button("▲", key=f"up_{sec['id']}_{q_idx}") and q_idx > 0:
                                sec["questions"][q_idx], sec["questions"][q_idx-1] = sec["questions"][q_idx-1], sec["questions"][q_idx]
                                st.rerun()
                        with btn_dn:
                            if st.button("▼", key=f"dn_{sec['id']}_{q_idx}") and q_idx < len(sec["questions"]) - 1:
                                sec["questions"][q_idx], sec["questions"][q_idx+1] = sec["questions"][q_idx+1], sec["questions"][q_idx]
                                st.rerun()

                        if st.button("❌ Remove", key=f"rem_q_{sec['id']}_{q_idx}"):
                            q_to_remove = q_idx

                        if st.button("💾 Save Bank", key=f"save_db_{sec['id']}_{q_idx}"):
                            q_save_payload = {
                                "subject": st.session_state.metadata["subject_title"],
                                "chapter": sec["title"],
                                "text": q["text"],
                                "is_mcq": sec["is_mcq"],
                                "options": q.get("options", {}),
                                "correct_answer": q.get("correct_answer", ""),
                                "marks": sec["marks_per_q"],
                                "difficulty": q.get("difficulty", "Medium"),
                                "blooms_level": q.get("blooms_level", "Understand")
                            }
                            save_question_to_db(q_save_payload)
                            st.success("Saved to DB!")

                    st.markdown("<hr style='border-color: rgba(255,255,255,0.05);'>", unsafe_allow_html=True)

                if q_to_remove is not None:
                    sec["questions"].pop(q_to_remove)
                    st.rerun()

                if st.button(f"➕ Add Question to {sec['title']}", key=f"add_q_{sec['id']}"):
                    sec["questions"].append({
                        "text": "New sample question prompt...",
                        "is_mcq": sec["is_mcq"],
                        "options": {"A": "Option 1", "B": "Option 2", "C": "Option 3", "D": "Option 4"},
                        "correct_answer": "A" if sec["is_mcq"] else "",
                        "difficulty": "Medium",
                        "blooms_level": "Understand"
                    })
                    st.rerun()

        if sec_to_remove is not None:
            st.session_state.sections.pop(sec_to_remove)
            st.rerun()

    with tab_parser:
        st.markdown("<div class='glass-card'>", unsafe_allow_html=True)
        st.subheader("📂 Upload & Parse Existing Sample Papers")
        uploaded_sample = st.file_uploader("Upload Sample Paper (.pdf or .docx)", type=["pdf", "docx"])
        if uploaded_sample:
            file_bytes = uploaded_sample.read()
            extracted_text = ""
            if uploaded_sample.name.endswith(".pdf"):
                extracted_text = extract_text_from_pdf(file_bytes)
            elif uploaded_sample.name.endswith(".docx"):
                extracted_text = extract_text_from_docx(file_bytes)

            st.text_area("Extracted Raw Document Text Preview", value=extracted_text, height=180)
            target_sec_index = st.selectbox(
                "Select Destination Section for Extracted Questions",
                options=list(range(len(st.session_state.sections))),
                format_func=lambda i: st.session_state.sections[i]["title"]
            )

            if st.button("⚡ Parse & Merge Questions into Selected Section"):
                parsed_qs = parse_raw_text_to_questions(extracted_text)
                if parsed_qs:
                    st.session_state.sections[target_sec_index]["questions"].extend(parsed_qs)
                    st.success(f"Successfully extracted and appended {len(parsed_qs)} questions!")
                else:
                    st.warning("No structured questions could be extracted.")
        st.markdown("</div>", unsafe_allow_html=True)

    with tab_bank:
        st.markdown("<div class='glass-card'>", unsafe_allow_html=True)
        st.subheader("🗄️ SQLite Persistent Question Bank")
        b_col1, b_col2 = st.columns(2)
        with b_col1:
            unique_subs = ["All"] + get_unique_subjects()
            filter_sub = st.selectbox("Filter by Subject", unique_subs)
        with b_col2:
            filter_diff = st.selectbox("Filter by Difficulty", ["All", "Easy", "Medium", "Hard"])

        bank_questions = get_filtered_questions(filter_sub, filter_diff)
        st.write(f"Found **{len(bank_questions)}** saved questions.")

        if bank_questions:
            for b_q in bank_questions:
                st.markdown(f"**[{b_q['subject']}]** {b_q['text']}")
                target_import_sec = st.selectbox(
                    "Import to Section",
                    options=list(range(len(st.session_state.sections))),
                    format_func=lambda i: st.session_state.sections[i]["title"],
                    key=f"imp_sec_{b_q['id']}"
                )
                if st.button(f"📥 Pull Question #{b_q['id']}", key=f"pull_{b_q['id']}"):
                    st.session_state.sections[target_import_sec]["questions"].append({
                        "text": b_q["text"], "is_mcq": b_q["is_mcq"], "options": b_q["options"],
                        "correct_answer": b_q["correct_answer"], "difficulty": b_q["difficulty"], "blooms_level": b_q["blooms_level"]
                    })
                    st.success("Imported!")
                st.markdown("<hr style='border-color: rgba(255,255,255,0.05);'>", unsafe_allow_html=True)
        st.markdown("</div>", unsafe_allow_html=True)

    with tab_randomizer:
        st.markdown("<div class='glass-card'>", unsafe_allow_html=True)
        st.subheader("🔀 Automated Multi-Set Generator (Set A / B / C)")
        num_sets = st.radio("Number of Randomized Sets", [2, 3], horizontal=True)

        if st.button("🎲 Generate Scrambled Exam Sets"):
            st.session_state.generated_sets = {}
            set_labels = ["Set A", "Set B", "Set C"][:num_sets]

            for label in set_labels:
                scrambled_sections = []
                for sec in st.session_state.sections:
                    sec_copy = {
                        "id": sec["id"], "title": sec["title"], "is_mcq": sec["is_mcq"],
                        "marks_per_q": sec["marks_per_q"], "questions": []
                    }
                    q_pool = list(sec["questions"])
                    random.shuffle(q_pool)

                    for q in q_pool:
                        q_copy = dict(q)
                        if sec["is_mcq"]:
                            orig_opts = list(q["options"].items())
                            random.shuffle(orig_opts)
                            new_opts, new_correct = {}, ""
                            keys = ["A", "B", "C", "D"]
                            for i, (old_k, val) in enumerate(orig_opts):
                                new_k = keys[i]
                                new_opts[new_k] = val
                                if old_k == q.get("correct_answer"):
                                    new_correct = new_k
                            q_copy["options"], q_copy["correct_answer"] = new_opts, new_correct
                        sec_copy["questions"].append(q_copy)
                    scrambled_sections.append(sec_copy)
                st.session_state.generated_sets[label] = scrambled_sections
            st.success("Generated randomized sets successfully!")

        if "generated_sets" in st.session_state:
            st.markdown("---")
            for set_name, set_sections in st.session_state.generated_sets.items():
                st.markdown(f"### 📄 {set_name}")
                set_meta = dict(st.session_state.metadata)
                set_meta["exam_name"] = f"{set_meta['exam_name']} - {set_name}"
                
                col_p, col_w = st.columns(2)
                with col_p:
                    st.download_button(f"⬇️ Download PDF ({set_name})", data=generate_pdf_paper(set_meta, set_sections), file_name=f"Exam_{set_name}.pdf", mime="application/pdf", key=f"dl_pdf_{set_name}")
                with col_w:
                    st.download_button(f"⬇️ Download Word ({set_name})", data=generate_docx_paper(set_meta, set_sections), file_name=f"Exam_{set_name}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", key=f"dl_docx_{set_name}")
        st.markdown("</div>", unsafe_allow_html=True)

    with tab_export:
        st.markdown("<div class='glass-card'>", unsafe_allow_html=True)
        st.subheader("🚀 High-Precision Document Generation Engine")
        ex_col1, ex_col2 = st.columns(2)

        with ex_col1:
            st.markdown("### 📄 Master PDF Export")
            st.download_button("📥 Download Master Exam Paper (PDF)", data=generate_pdf_paper(st.session_state.metadata, st.session_state.sections, False), file_name="Master_Examination_Paper.pdf", mime="application/pdf", use_container_width=True)
            st.markdown("<br>", unsafe_allow_html=True)
            st.download_button("🔑 Download Solution Key (PDF)", data=generate_pdf_paper(st.session_state.metadata, st.session_state.sections, True), file_name="Master_Answer_Key.pdf", mime="application/pdf", use_container_width=True)

        with ex_col2:
            st.markdown("### 📝 Master Word Export")
            st.download_button("📥 Download Master Exam Paper (.docx)", data=generate_docx_paper(st.session_state.metadata, st.session_state.sections, False), file_name="Master_Examination_Paper.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
            st.markdown("<br>", unsafe_allow_html=True)
            st.download_button("🔑 Download Solution Key (.docx)", data=generate_docx_paper(st.session_state.metadata, st.session_state.sections, True), file_name="Master_Answer_Key.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
        st.markdown("</div>", unsafe_allow_html=True)

if __name__ == "__main__":
    main()